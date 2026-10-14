from datetime import timedelta

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from .models import ParkingArea, ParkingFloor, ParkingRecord, ParkingSlot, Vehicle


class ParkingDateAndFloorTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='operator', password='pass12345')
        self.client.login(username='operator', password='pass12345')
        self.floor_two_slot = ParkingSlot.objects.get(number='F2-01')
        self.ground_slot = ParkingSlot.objects.get(number='G-01')

    def open_management(self):
        """Complete the required fresh staff authentication for management tests."""
        response = self.client.post(reverse('parking_management'), {
            'action': 'management_reauth', 'username': 'operator', 'password': 'pass12345',
        })
        self.assertEqual(response.status_code, 302)
        return self.client.get(response['Location'])

    def test_entry_uses_current_datetime_and_assigns_floor_slot(self):
        started_at = timezone.now()
        response = self.client.post(reverse('vehicle_entry'), {
            'vehicle_number': 'GJ 01 AB 1234',
            'state_code': 'GJ',
            'vehicle_type': 'Car',
            'plate_type': 'white_black',
            'parking_area': self.floor_two_slot.parking_floor.parking_area_id,
            'parking_floor': self.floor_two_slot.parking_floor_id,
            'parking_slot': self.floor_two_slot.id,
            'entry_time': '2026-08-17T10:00',
            'owner_name': 'Test Owner',
            'mobile': '9876543210',
        })

        self.assertRedirects(response, reverse('dashboard'))
        record = ParkingRecord.objects.get()
        self.assertEqual(record.parking_number, 'F2-01')
        self.assertEqual(record.floor_name, 'Floor 2')
        self.assertGreaterEqual(record.entry_time, started_at)
        self.assertLess(record.entry_time, timezone.now() + timedelta(seconds=1))

    def test_exit_uses_current_datetime_for_fee(self):
        self.client.post(reverse('vehicle_entry'), {
            'vehicle_number': 'GJ 01 AB 1234', 'state_code': 'GJ', 'vehicle_type': 'Car',
            'plate_type': 'white_black', 'parking_area': self.ground_slot.parking_floor.parking_area_id,
            'parking_floor': self.ground_slot.parking_floor_id, 'parking_slot': self.ground_slot.id,
            'entry_time': '2026-08-17T10:00', 'owner_name': 'Test Owner', 'mobile': '9876543210',
        })
        record = ParkingRecord.objects.get()
        record.entry_time = timezone.now() - timedelta(hours=2, minutes=15)
        record.save(update_fields=('entry_time',))
        checkout_started = timezone.now()
        response = self.client.post(reverse('vehicle_exit'), {
            'vehicle_number': 'GJ 01 AB 1234', 'exit_time': '2026-08-17T12:15',
        })

        record.refresh_from_db()
        self.assertRedirects(response, reverse('receipt', args=[record.id]))
        self.assertEqual(record.status, 'Exited')
        self.assertGreaterEqual(record.exit_time, checkout_started)
        self.assertEqual(record.parking_fee, 90)

    def test_bh_series_forces_bh_state_code(self):
        response = self.client.post(reverse('vehicle_entry'), {
            'vehicle_number': 'BH 01 AA 0001', 'state_code': 'XX', 'bh_series': 'on',
            'vehicle_type': 'Car', 'plate_type': 'bh_series',
            'parking_area': self.ground_slot.parking_floor.parking_area_id,
            'parking_floor': self.ground_slot.parking_floor_id,
            'parking_slot': self.ground_slot.id,
            'entry_time': '2026-08-17T10:00', 'owner_name': 'BH Owner', 'mobile': '9876543210',
        })

        self.assertRedirects(response, reverse('dashboard'))
        self.assertEqual(ParkingRecord.objects.get().vehicle.state_code, 'BH')

    def test_dashboard_renders_a_floor_by_floor_parking_map(self):
        response = self.client.get(reverse('dashboard'))

        self.assertContains(response, 'Live Parking Layout')
        self.assertContains(response, 'Ground Floor')
        self.assertContains(response, 'Floor 3')

    def test_staff_can_open_custom_parking_management_page(self):
        self.user.is_staff = True
        self.user.save(update_fields=('is_staff',))

        response = self.open_management()

        self.assertContains(response, 'Parking Management')
        self.assertContains(response, 'Parking Area')
        self.assertNotContains(response, 'Django administration')

    def test_staff_can_create_an_area_then_a_floor_with_generated_slots(self):
        self.user.is_staff = True
        self.user.save(update_fields=('is_staff',))

        self.open_management()
        response = self.client.post(reverse('parking_management'), {
            'action': 'add_area', 'name': 'Visitor Parking', 'address': 'North gate', 'total_slots': 100,
        })
        self.assertRedirects(response, reverse('parking_management'))
        area = ParkingArea.objects.get(name='Visitor Parking')

        self.open_management()
        response = self.client.post(reverse('parking_management'), {
            'action': 'add_floor', 'parking_area': area.id, 'name': 'Ground Floor',
            'slot_prefix': 'G', 'slot_count': 30, 'sort_order': 1,
        })
        self.assertRedirects(response, reverse('parking_management'))
        floor = ParkingFloor.objects.get(parking_area=area, name='Ground Floor')
        self.assertEqual(floor.slots.count(), 30)
        self.assertTrue(ParkingSlot.objects.filter(parking_floor=floor, number='G-001').exists())
        self.assertTrue(ParkingSlot.objects.filter(parking_floor=floor, number='G-030').exists())

    def test_staff_can_save_the_layout_in_one_action(self):
        self.user.is_staff = True
        self.user.save(update_fields=('is_staff',))
        area = ParkingArea.objects.get(name='Main Parking')
        data = {
            'action': 'save_layout',
            f'area_name_{area.id}': area.name,
            f'area_address_{area.id}': area.address,
            f'area_payment_qr_data_{area.id}': area.payment_qr_data,
            f'area_total_slots_{area.id}': area.total_slots,
        }
        for floor in area.floors.all():
            data.update({
                f'floor_name_{floor.id}': floor.name,
                f'floor_slot_prefix_{floor.id}': floor.slot_prefix,
                f'floor_slot_count_{floor.id}': floor.slots.count(),
                f'floor_sort_order_{floor.id}': floor.sort_order,
            })
            for slot in floor.slots.all():
                data.update({
                    f'slot_number_{slot.id}': slot.number,
                    f'slot_sort_order_{slot.id}': slot.sort_order,
                })

        data[f'floor_maintenance_{self.ground_slot.parking_floor_id}'] = 'on'
        self.open_management()
        response = self.client.post(reverse('parking_management'), data)

        self.assertRedirects(response, reverse('parking_management'))
        self.ground_slot.parking_floor.refresh_from_db()
        self.assertTrue(self.ground_slot.parking_floor.is_under_maintenance)

    def test_slot_under_maintenance_reduces_available_not_total(self):
        initial_total = ParkingSlot.objects.count()
        response_before = self.client.get(reverse('dashboard'))
        initial_available = response_before.context['available_slots']

        # Put a slot under maintenance
        self.floor_two_slot.is_under_maintenance = True
        self.floor_two_slot.save()

        response_after = self.client.get(reverse('dashboard'))
        # Total slots must NOT change when a slot is under maintenance
        self.assertEqual(response_after.context['total_slots'], initial_total)
        # Available slots MUST reduce
        self.assertEqual(response_after.context['available_slots'], initial_available - 1)
        # Dashboard renders the slot with maintenance class
        self.assertContains(response_after, 'is-maintenance')

        # Deleting a slot DOES reduce total_slots
        self.floor_two_slot.delete()
        response_deleted = self.client.get(reverse('dashboard'))
        self.assertEqual(response_deleted.context['total_slots'], initial_total - 1)


class AnalyticsDashboardAndReportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='operator', password='pass12345')
        self.client.login(username='operator', password='pass12345')
        self.ground_slot = ParkingSlot.objects.get(number='G-01')

        # Create sample vehicle and parking records
        v1 = Vehicle.objects.create(
            vehicle_number='MH 02 CZ 9999',
            owner_name='Aditya Roy',
            mobile='9876543210',
            vehicle_type='Car',
            plate_type='white_black',
            state_code='MH'
        )
        now = timezone.now()
        ParkingRecord.objects.create(
            vehicle=v1,
            parking_slot=self.ground_slot,
            slot_number=1,
            entry_time=now - timedelta(hours=3),
            exit_time=now,
            total_hours=3.0,
            parking_fee=90.00,
            status='Exited'
        )

    def test_analytics_dashboard_renders(self):
        response = self.client.get(reverse('analytics'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Analytics & Intelligence Hub')
        self.assertContains(response, 'Total Revenue')
        self.assertContains(response, 'Peak Parking Times')
        self.assertIn('kpis', response.context)
        self.assertEqual(response.context['kpis']['total_records'], 1)
        self.assertEqual(response.context['kpis']['total_revenue'], 90.0)

    def test_analytics_filter_by_vehicle_type(self):
        response = self.client.get(reverse('analytics') + '?vehicle_type=Car')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['kpis']['total_records'], 1)

        response_empty = self.client.get(reverse('analytics') + '?vehicle_type=Bus')
        self.assertEqual(response_empty.status_code, 200)
        self.assertEqual(response_empty.context['kpis']['total_records'], 0)

    def test_export_analytics_excel(self):
        response = self.client.get(reverse('export_analytics_excel'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment; filename="ParkNest_Analytics_Report_', response['Content-Disposition'])

        # Verify openpyxl can load the generated Excel workbook
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn('Parking Records', wb.sheetnames)
        self.assertIn('Executive Summary', wb.sheetnames)
        ws = wb['Parking Records']
        self.assertIn('PARK NEST — COMPREHENSIVE PARKING AUDIT REPORT', ws['A1'].value)

    def test_export_analytics_excel_custom_columns(self):
        response = self.client.get(reverse('export_analytics_excel'), {
            'columns': ['vehicle_number', 'fee', 'status'],
            'include_summary': '0'
        })
        self.assertEqual(response.status_code, 200)
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn('Parking Records', wb.sheetnames)
        self.assertNotIn('Executive Summary', wb.sheetnames)

    def test_seed_and_clear_demo_analytics(self):
        seed_resp = self.client.get(reverse('seed_demo_analytics'))
        self.assertRedirects(seed_resp, reverse('analytics'))
        self.assertGreater(ParkingRecord.objects.count(), 10)

        clear_resp = self.client.get(reverse('clear_demo_analytics'))
        self.assertRedirects(clear_resp, reverse('analytics'))


